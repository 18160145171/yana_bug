#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import csv
import datetime as dt
import html
import io
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REQUIRED_COLUMNS = []
OPTIONAL_COLUMNS = ["标题", "任务ID", "执行者", "父任务", "任务状态", "优先级", "参与者"]

COLUMN_ALIASES = {
    "标题": ["标题", "bug标题", "缺陷标题", "问题标题", "title", "summary", "主题"],
    "任务ID": ["任务id", "任务编号", "单号", "id", "bugid", "缺陷id", "ticketid"],
    "执行者": ["执行者", "处理人", "责任人", "指派给", "assignee", "owner", "执行人"],
    "父任务": ["父任务", "父任务id", "上级任务", "所属任务", "parent", "parentid"],
    "任务状态": ["任务状态", "状态", "bug状态", "缺陷状态", "当前状态", "status"],
    "优先级": ["优先级", "严重程度", "等级", "priority", "severity"],
    "参与者": ["参与者", "参与人", "抄送人", "协作者", "关注人", "watchers", "cc"],
}

UNCLOSED_STATUS_ORDER = ["重新打开", "待处理", "开发中", "暂不处理"]
CLOSED_STATUS_ORDER = ["已解决", "已完成", "开发完成"]

STATUS_GROUP_INDEX = {s: 0 for s in UNCLOSED_STATUS_ORDER}
STATUS_GROUP_INDEX.update({s: 1 for s in CLOSED_STATUS_ORDER})

ALL_STATUS_ORDER = UNCLOSED_STATUS_ORDER + CLOSED_STATUS_ORDER
STATUS_INDEX = {status: idx for idx, status in enumerate(ALL_STATUS_ORDER)}

HIGH_PRIORITY_KEYWORDS = ["高", "紧急", "严重", "critical", "p0", "p1", "blocker"]
UI_KEYWORDS = ["ui", "页面", "显示", "样式", "布局", "交互", "按钮", "弹窗", "对齐", "颜色", "字体"]
MODEL_MAX_ROWS = 500
DEFAULT_API_ENDPOINT = "/v1/chat/completions"

DOCX_BODY_FONT = "宋体"
DOCX_HEADING_FONT = "黑体"
DOCX_ACCENT = "4F81BD"
DOCX_BORDER = "D9D9D9"
DOCX_HEADER_FILL = "D9EAF7"
DOCX_ALT_ROW_FILL = "F7F9FC"

DEFAULT_REPORT_PROMPT = """请以优秀测试软件工程师和测试经理的视角，重新分析输入的 Bug 数据，生成正式、客观、简洁、可执行的质量验收与缺陷分析结论。
严格以输入数据为依据，不臆造版本、环境、复现步骤、责任人或修复状态；输入缺失的信息请明确标注“数据未提供”。
重点分析 Bug 总量、有效记录、优先级分布、功能模块分布、执行人分布、任务状态、已解决率、未关闭问题和高优先级未闭环问题。
识别缺陷集中模块、发布风险、数据质量问题（例如标题不完整、状态与优先级异常或关键信息缺失），并给出有依据的发布建议和后续行动。
结论要说明判断依据；行动建议要具体到责任角色、优先级或验证动作。请使用中文。"""

MODULE_RULES = [
    ("去水印/去字幕", ["去水印", "去字幕", "logo"]),
    ("数据埋点", ["埋点"]),
    ("多语言适配", ["多语言"]),
    ("视频翻译", ["视频翻译"]),
    ("AI字幕", ["ai字幕", "ai 字幕", "智能字幕"]),
    ("设置页", ["设置页", "设置页面"]),
    ("编辑页", ["编辑页", "编辑页面"]),
]


def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize_header_name(name):
    raw = normalize_text(name).lower()
    # 移除空格、下划线和常见标点，提升列名兼容性
    return re.sub(r"[\s\-_（）()\[\]【】:：]+", "", raw)


def build_header_mapping(headers):
    normalized_to_original = {normalize_header_name(h): h for h in headers}
    mapping = {}

    for canonical, alias_list in COLUMN_ALIASES.items():
        found = None
        for alias in alias_list:
            alias_norm = normalize_header_name(alias)
            if alias_norm in normalized_to_original:
                found = normalized_to_original[alias_norm]
                break

        if not found:
            # 包含匹配，兼容如“状态(当前)”这类表头
            for normalized_header, original_header in normalized_to_original.items():
                if any(normalize_header_name(alias) in normalized_header for alias in alias_list):
                    found = original_header
                    break

        if found:
            mapping[canonical] = found

    return mapping


def split_people(text):
    value = normalize_text(text)
    if not value:
        return []
    parts = re.split(r"[，,、;/\\\s]+", value)
    return [p.strip() for p in parts if p.strip()]


def is_high_priority(priority):
    p = normalize_text(priority).lower()
    if not p:
        return False
    return any(token in p for token in HIGH_PRIORITY_KEYWORDS)


def classify_module(title):
    text = normalize_text(title).lower()
    for module_name, keywords in MODULE_RULES:
        if any(k.lower() in text for k in keywords):
            return module_name

    if "设置" in text:
        return "设置页"
    if "编辑" in text:
        return "编辑页"
    if "字幕" in text and "ai" in text:
        return "AI字幕"
    if "翻译" in text and "视频" in text:
        return "视频翻译"
    return "其他模块"


def issue_nature_by_titles(titles):
    if not titles:
        return "功能逻辑问题"
    ui_hits = 0
    for title in titles:
        t = normalize_text(title).lower()
        if any(k in t for k in UI_KEYWORDS):
            ui_hits += 1
    ratio = ui_hits / len(titles)
    return "UI交互问题" if ratio >= 0.5 else "功能逻辑问题"


def parse_task_id(task_id):
    text = normalize_text(task_id)
    nums = re.findall(r"\d+", text)
    if nums:
        return int(nums[-1]), text
    return sys.maxsize, text


def load_rows(input_file):
    ext = Path(input_file).suffix.lower()
    if ext == ".csv":
        with open(input_file, "r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            rows = [dict(r) for r in reader]
            headers = reader.fieldnames or []
        return harmonize_columns(headers, rows)

    if ext in [".xlsx", ".xls"]:
        try:
            import pandas as pd  # pylint: disable=import-outside-toplevel
        except ImportError as exc:
            raise RuntimeError(
                "读取 Excel 需要 pandas/openpyxl。请先执行：pip install pandas openpyxl"
            ) from exc
        df = pd.read_excel(input_file, dtype=str).fillna("")
        headers = list(df.columns)
        rows = df.to_dict(orient="records")
        return harmonize_columns(headers, rows)

    raise RuntimeError("仅支持 .csv / .xlsx / .xls 文件。")


def validate_columns(headers):
    if not headers:
        raise RuntimeError("未识别到可用字段，请检查表头行是否存在。")


def harmonize_columns(headers, rows):
    mapping = build_header_mapping(headers)

    normalized_rows = []
    for row in rows:
        normalized_row = {}
        for canonical in OPTIONAL_COLUMNS:
            source_col = mapping.get(canonical)
            normalized_row[canonical] = normalize_text(row.get(source_col, "")) if source_col else ""

        # 参与者缺失时，默认沿用执行者
        if not normalized_row.get("参与者"):
            normalized_row["参与者"] = normalized_row.get("执行者", "")

        normalized_rows.append(normalized_row)

    detected_headers = list(mapping.keys())
    return detected_headers, normalized_rows


def density_weight(count):
    if count > 5:
        return 3
    if 2 <= count <= 5:
        return 2
    return 1


def format_percent(numerator, denominator):
    if denominator == 0:
        return "0.00%"
    return f"{(numerator / denominator) * 100:.2f}%"


def sort_rows_for_appendix(rows):
    def sort_key(row):
        status = normalize_text(row.get("任务状态"))
        group = STATUS_GROUP_INDEX.get(status, 2)
        status_idx = STATUS_INDEX.get(status, 999)
        high_rank = 0 if is_high_priority(row.get("优先级")) else 1
        task_num, task_raw = parse_task_id(row.get("任务ID"))
        return (group, status_idx, high_rank, task_num, task_raw)

    return sorted(rows, key=sort_key)


def ensure_f_drive_or_desktop():
    f_drive = Path("F:/")
    if f_drive.exists():
        return f_drive
    desktop = Path.home() / "Desktop"
    return desktop


def resolve_api_endpoint(api_base_url, api_endpoint):
    base = normalize_text(api_base_url).rstrip("/")
    if not base:
        raise RuntimeError("API Base URL 不能为空。")

    configured_endpoint = normalize_text(api_endpoint) or DEFAULT_API_ENDPOINT
    if configured_endpoint.lower().startswith(("http://", "https://")):
        return configured_endpoint.rstrip("/")

    path = configured_endpoint if configured_endpoint.startswith("/") else f"/{configured_endpoint}"
    if path.lower().startswith("/v1/") and base.lower().endswith("/v1"):
        path = path[3:]
    if base.lower().endswith(path.lower()):
        return base
    return f"{base}{path}"


def _is_responses_endpoint(endpoint):
    path = urllib.parse.urlparse(endpoint).path.rstrip("/").lower()
    return path.endswith("/responses")


def _alternate_compatible_endpoint(endpoint):
    parsed = urllib.parse.urlsplit(endpoint)
    path = parsed.path.rstrip("/")
    if not path.lower().endswith("/responses"):
        return ""
    alternate_path = f"{path[:-len('/responses')]}/chat/completions"
    return urllib.parse.urlunsplit(
        (parsed.scheme, parsed.netloc, alternate_path, parsed.query, parsed.fragment)
    )


def _is_cloudflare_1010(detail):
    lowered = normalize_text(detail).lower()
    return "1010" in lowered and (
        "cloudflare" in lowered
        or "error code" in lowered
        or "error 1010" in lowered
    )


def _messages_to_responses_input(messages):
    return [
        {
            "role": normalize_text(message.get("role")) or "user",
            "content": [
                {
                    "type": "input_text",
                    "text": normalize_text(message.get("content")),
                }
            ],
        }
        for message in messages
    ]


def _extract_response_text(data):
    if not isinstance(data, dict):
        raise RuntimeError("模型返回格式异常：顶层结果不是 JSON 对象。")

    output_text = data.get("output_text")
    if isinstance(output_text, str) and output_text.strip():
        return output_text.strip()

    choices = data.get("choices")
    if isinstance(choices, list) and choices:
        message = choices[0].get("message", {})
        content = message.get("content") if isinstance(message, dict) else None
        if isinstance(content, str):
            return content.strip()
        if isinstance(content, list):
            text_parts = [
                item.get("text", "")
                for item in content
                if isinstance(item, dict) and isinstance(item.get("text"), str)
            ]
            return "".join(text_parts).strip()

    output = data.get("output")
    text_parts = []
    if isinstance(output, list):
        for item in output:
            if not isinstance(item, dict):
                continue
            content = item.get("content")
            if isinstance(content, str):
                text_parts.append(content)
            elif isinstance(content, list):
                for part in content:
                    if isinstance(part, dict) and isinstance(part.get("text"), str):
                        text_parts.append(part["text"])
    if text_parts:
        return "".join(text_parts).strip()

    raise RuntimeError(f"模型返回格式异常：未找到文本内容。原始返回：{json.dumps(data, ensure_ascii=False)[:500]}")


class _ApiEndpointError(RuntimeError):
    def __init__(self, status_code, endpoint, detail):
        self.status_code = status_code
        self.endpoint = endpoint
        self.detail = detail
        super().__init__(self._build_message())

    def _build_message(self):
        if self.status_code == 401:
            return (
                "模型接口返回 401（鉴权失败）。请检查："
                "1) API Key 是否有效且完整；"
                f"2) 接口地址是否正确（当前请求：{self.endpoint}）；"
                "3) 模型标识是否与服务商配置一致。"
                f" 原始返回：{self.detail[:300]}"
            )
        if self.status_code == 403:
            return (
                f"模型接口返回 403（服务端拒绝请求），当前请求地址：{self.endpoint}。"
                "请检查接口端点是否被服务商或 Cloudflare/WAF 拦截。"
                f" 原始返回：{self.detail[:500]}"
            )
        return (
            f"模型接口错误 HTTP {self.status_code}，"
            f"当前请求地址：{self.endpoint}：{self.detail[:500]}"
        )


def _request_openai_compatible_endpoint(
    endpoint,
    api_key,
    model_name,
    messages,
    timeout_sec,
):
    if _is_responses_endpoint(endpoint):
        payload = {
            "model": model_name,
            "input": _messages_to_responses_input(messages),
        }
    else:
        payload = {
            "model": model_name,
            "messages": messages,
            "temperature": 0.3,
        }
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(
        endpoint,
        data=body,
        headers={
            "Content-Type": "application/json",
            "Accept": "application/json",
            "Authorization": f"Bearer {api_key}",
            # Some OpenAI-compatible gateways reject Python-urllib's default signature at the WAF layer.
            "User-Agent": "OpenAI/Python",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout_sec) as resp:
            raw = resp.read().decode("utf-8", errors="replace")
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise _ApiEndpointError(exc.code, endpoint, detail) from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"模型接口连接失败（{endpoint}）：{exc.reason}") from exc

    try:
        data = json.loads(raw)
        return _extract_response_text(data)
    except Exception as exc:  # pylint: disable=broad-except
        if isinstance(exc, RuntimeError):
            raise
        raise RuntimeError(f"模型返回格式异常: {raw[:500]}") from exc


def call_openai_compatible_chat(
    api_base_url,
    api_key,
    model_name,
    messages,
    timeout_sec=45,
    api_endpoint=DEFAULT_API_ENDPOINT,
):
    base = normalize_text(api_base_url).rstrip("/")
    endpoint = resolve_api_endpoint(api_base_url, api_endpoint)

    normalized_model = normalize_text(model_name)
    if "open.bigmodel.cn" in base.lower():
        alias_map = {
            "智谱": "glm-4-plus",
            "zhipu": "glm-4-plus",
            "zhipuai": "glm-4-plus",
            "glm4": "glm-4-plus",
        }
        normalized_model = alias_map.get(normalized_model.lower(), normalized_model) if normalized_model else "glm-4-plus"
        if normalized_model == "智谱":
            normalized_model = "glm-4-plus"
        if "." not in normalize_text(api_key):
            raise RuntimeError(
                "检测到智谱接口，但 API Key 格式疑似不正确。"
                "通常应为 'id.secret' 形式，请从智谱控制台重新复制。"
            )

    try:
        return _request_openai_compatible_endpoint(
            endpoint=endpoint,
            api_key=api_key,
            model_name=normalized_model,
            messages=messages,
            timeout_sec=timeout_sec,
        )
    except _ApiEndpointError as primary_error:
        alternate_endpoint = _alternate_compatible_endpoint(endpoint)
        if not (
            primary_error.status_code == 403
            and _is_cloudflare_1010(primary_error.detail)
            and alternate_endpoint
        ):
            raise RuntimeError(str(primary_error)) from primary_error

        try:
            return _request_openai_compatible_endpoint(
                endpoint=alternate_endpoint,
                api_key=api_key,
                model_name=normalized_model,
                messages=messages,
                timeout_sec=timeout_sec,
            )
        except Exception as alternate_error:  # pylint: disable=broad-except
            raise RuntimeError(
                f"主端点 {endpoint} 返回 Cloudflare 1010（403），"
                f"已自动切换到兼容端点 {alternate_endpoint}，但仍调用失败："
                f"{alternate_error}"
            ) from alternate_error


def build_ai_insight(project_name, total, resolved_count, high_count, module_analysis, top1_module, top1_nature, rows, llm_config):
    unresolved = []
    for row in rows:
        status = normalize_text(row.get("任务状态"))
        if status in UNCLOSED_STATUS_ORDER:
            unresolved.append(
                {
                    "title": normalize_text(row.get("标题")),
                    "status": status,
                    "priority": normalize_text(row.get("优先级")),
                    "owner": normalize_text(row.get("执行者")),
                }
            )
    unresolved = unresolved[:15]

    module_stats = [
        {
            "module": item["module"],
            "count": item["count"],
            "density": round(item["density"], 2),
        }
        for item in module_analysis[:8]
    ]

    model_rows = sort_rows_for_appendix(rows)[:MODEL_MAX_ROWS]
    context = {
        "project_name": project_name,
        "total_bug_count": total,
        "resolved_count": resolved_count,
        "high_priority_count": high_count,
        "resolved_rate": format_percent(resolved_count, total),
        "high_priority_rate": format_percent(high_count, total),
        "top1_module": top1_module,
        "top1_issue_nature": top1_nature,
        "module_stats": module_stats,
        "sample_unresolved_bugs": unresolved,
        "bug_rows": model_rows,
        "bug_rows_total": total,
        "bug_rows_included": len(model_rows),
        "bug_rows_truncated": total > len(model_rows),
    }

    report_prompt = normalize_text(llm_config.get("report_prompt")) or DEFAULT_REPORT_PROMPT
    messages = [
        {
            "role": "system",
            "content": (
                "你是优秀测试软件工程师和资深测试经理。"
                "请根据用户提供的报告生成规则，重新分析 Bug 明细和统计信息，生成客观、可执行的发布评估。"
                "只输出 JSON，不要输出 Markdown 代码块。"
            ),
        },
        {
            "role": "user",
            "content": (
                "【报告生成规则提示词】\n"
                f"{report_prompt}\n\n"
                "【输出格式约束】\n"
                "请严格返回 JSON 对象，字段必须包含："
                "executive_summary, key_risks(数组), release_decision, actions(数组)。"
                "其中 release_decision 只能是 GO / NO-GO / CONDITIONAL-GO。"
                "每个 action 最多 30 字。不得输出 JSON 以外的内容。"
                "这些字段会直接写入报告的风险结论、发布建议和改进措施，"
                "请按报告生成规则重新判断，不要照抄固定模板，也不要补写输入数据中不存在的事实。"
                "以下是本次上传文件经过字段标准化后的统计和 Bug 明细；"
                "如果 bug_rows_truncated 为 true，说明明细过多，仅展示优先级更高或未关闭的前 500 条，"
                "请结合总数和统计字段进行判断：\n"
                + json.dumps(context, ensure_ascii=False)
            ),
        },
    ]

    content = call_openai_compatible_chat(
        api_base_url=llm_config.get("api_base_url", ""),
        api_key=llm_config.get("api_key", ""),
        model_name=llm_config.get("model_name", ""),
        messages=messages,
        api_endpoint=llm_config.get("api_endpoint", DEFAULT_API_ENDPOINT),
    )

    cleaned = content.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
        cleaned = re.sub(r"\s*```$", "", cleaned)

    try:
        result = json.loads(cleaned)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"模型返回不是合法 JSON：{cleaned[:500]}") from exc
    if not isinstance(result, dict):
        raise RuntimeError("模型返回格式异常：顶层结果必须是 JSON 对象。")

    required_fields = {"executive_summary", "key_risks", "release_decision", "actions"}
    missing_fields = sorted(required_fields - set(result))
    if missing_fields:
        raise RuntimeError(f"模型返回缺少字段：{', '.join(missing_fields)}")

    summary = normalize_text(result.get("executive_summary"))
    decision = normalize_text(result.get("release_decision")).upper()
    risks = result.get("key_risks", [])
    actions = result.get("actions", [])

    if decision not in {"GO", "NO-GO", "CONDITIONAL-GO"}:
        raise RuntimeError(
            "模型返回的 release_decision 无效，必须是 GO / NO-GO / CONDITIONAL-GO。"
        )
    if not summary:
        raise RuntimeError("模型返回的 executive_summary 为空。")
    if not isinstance(risks, list):
        raise RuntimeError("模型返回的 key_risks 必须是数组。")
    if not isinstance(actions, list):
        raise RuntimeError("模型返回的 actions 必须是数组。")

    return {
        "summary": summary,
        "decision": decision,
        "risks": [normalize_text(r) for r in risks if normalize_text(r)],
        "actions": [normalize_text(a) for a in actions if normalize_text(a)],
    }


def build_html(project_name, rows, llm_config=None):
    total = len(rows)
    resolved_count = sum(1 for r in rows if normalize_text(r.get("任务状态")) == "已解决")
    high_count = sum(1 for r in rows if is_high_priority(r.get("优先级")))

    participants = set()
    module_counter = Counter()
    module_titles = defaultdict(list)

    for row in rows:
        for person in split_people(row.get("执行者")):
            participants.add(person)
        title = normalize_text(row.get("标题"))
        module = classify_module(title)
        module_counter[module] += 1
        module_titles[module].append(title)

    module_analysis = []
    for module_name, count in module_counter.items():
        weight = density_weight(count)
        density = count / weight
        module_analysis.append(
            {
                "module": module_name,
                "count": count,
                "weight": weight,
                "density": density,
            }
        )
    module_analysis.sort(key=lambda x: (-x["count"], x["module"]))

    top1_module = module_analysis[0]["module"] if module_analysis else "无"
    top1_nature = issue_nature_by_titles(module_titles.get(top1_module, []))

    sorted_rows = sort_rows_for_appendix(rows)
    people_text = "、".join(sorted(participants)) if participants else "无"

    summary_box = (
        f"Bug总数：<b>{total}</b>；"
        f"解决率（已解决/总数）：<b>{format_percent(resolved_count, total)}</b>；"
        f"高优先级占比：<b>{format_percent(high_count, total)}</b>；"
        f"参与人员：<b>{html.escape(people_text)}</b>"
    )

    top_risk = "当前未发现明显风险。"
    if total > 0 and module_analysis:
        top_risk = (
            f"当前缺陷主要集中在「{html.escape(top1_module)}」模块，"
            f"Top1 模块倾向于 <b>{top1_nature}</b>，建议优先投入专项回归资源。"
        )

    module_rows_html = "".join(
        [
            (
                "<tr>"
                f"<td>{idx + 1}</td>"
                f"<td>{html.escape(item['module'])}</td>"
                f"<td>{item['count']}</td>"
                f"<td>{item['weight']}</td>"
                f"<td>{item['density']:.2f}</td>"
                "</tr>"
            )
            for idx, item in enumerate(module_analysis)
        ]
    )

    if not module_rows_html:
        module_rows_html = "<tr><td colspan='5'>暂无数据</td></tr>"

    improvements = [
        (
            "Top模块专项治理",
            f"针对「{html.escape(top1_module)}」建立专项缺陷清单，按 {top1_nature} 维度拆分并逐项闭环。",
        ),
        (
            "高优先级前置拦截",
            "将高优先级问题纳入每日跟踪，新增提测门禁：高优先级未清零禁止发布。",
        ),
        (
            "回归用例增强",
            "沉淀高频问题场景为自动化回归用例，覆盖跨页面、跨状态、跨语言主流程。",
        ),
    ]

    improvements_html = "".join(
        [
            (
                "<div class='improvement-item'>"
                f"<span class='improvement-title'>{html.escape(title)}</span>"
                f"<span>{html.escape(content)}</span>"
                "</div>"
            )
            for title, content in improvements
        ]
    )

    appendix_rows_html = ""
    for row in sorted_rows:
        status = normalize_text(row.get("任务状态"))
        cls = "status-unclosed" if status in UNCLOSED_STATUS_ORDER else "status-closed"
        appendix_rows_html += (
            "<tr>"
            f"<td>{html.escape(normalize_text(row.get('标题')))}</td>"
            f"<td>{html.escape(normalize_text(row.get('任务ID')))}</td>"
            f"<td>{html.escape(normalize_text(row.get('执行者')))}</td>"
            f"<td>{html.escape(normalize_text(row.get('父任务')))}</td>"
            f"<td class='{cls}'>{html.escape(status)}</td>"
            f"<td>{html.escape(normalize_text(row.get('优先级')))}</td>"
            "</tr>"
        )
    if not appendix_rows_html:
        appendix_rows_html = "<tr><td colspan='6'>暂无缺陷数据</td></tr>"

    ai_analysis = None
    if llm_config and llm_config.get("enabled"):
        # AI 模式必须真实完成接口调用并解析模型结果，失败时不生成模板报告。
        try:
            ai_analysis = build_ai_insight(
                project_name=project_name,
                total=total,
                resolved_count=resolved_count,
                high_count=high_count,
                module_analysis=module_analysis,
                top1_module=top1_module,
                top1_nature=top1_nature,
                rows=rows,
                llm_config=llm_config,
            )
        except Exception as exc:  # pylint: disable=broad-except
            raise RuntimeError(f"AI 增强调用失败，报告未生成：{exc}") from exc

    module_interpretation_html = (
        "<div class='summary-box'>"
        f"Top1模块：<b>{html.escape(top1_module)}</b>；"
        f"动态深度解读：该模块问题主要表现为 <b>{html.escape(top1_nature)}</b>。"
        "</div>"
    )
    ai_section_html = ""
    if ai_analysis:
        ai_risks = ai_analysis["risks"] or ["模型未识别出需要单独列出的关键风险。"]
        ai_actions = ai_analysis["actions"] or ["模型未返回后续行动，请根据数据继续确认。"]
        risk_lines = "".join(
            f"<li>{html.escape(item)}</li>" for item in ai_risks
        )
        action_lines = "".join(
            f"<li>{html.escape(item)}</li>" for item in ai_actions
        )

        # AI 成功后，报告叙述部分直接使用模型结果，不再展示固定模板结论。
        top_risk = (
            "<div class='summary-box'>"
            f"<b>AI发布建议：</b>{html.escape(ai_analysis['decision'])}<br>"
            f"<b>AI结论摘要：</b>{html.escape(ai_analysis['summary'])}"
            "</div>"
            "<div class='risk-box'><b>AI关键风险：</b><ul>"
            f"{risk_lines}</ul></div>"
        )
        module_interpretation_html = (
            "<div class='summary-box'>"
            "<b>AI重新分析结论：</b>"
            f"{html.escape(ai_analysis['summary'])}"
            "</div>"
        )
        improvements_html = "".join(
            (
                "<div class='improvement-item'>"
                f"<span class='improvement-title'>AI行动建议 {idx}</span>"
                f"<span>{html.escape(action)}</span>"
                "</div>"
            )
            for idx, action in enumerate(ai_actions, start=1)
        )
        ai_section_html = (
            "<h2>五、 AI 接口分析结果</h2>"
            "<div><b>模型返回的后续行动：</b><ul>"
            f"{action_lines}</ul></div>"
        )

    html_doc = f"""<html xmlns:o='urn:schemas-microsoft-com:office:office' xmlns:w='urn:schemas-microsoft-com:office:word' xmlns='http://www.w3.org/TR/REC-html40'>
<head>
    <meta charset="utf-8">
    <style>
        body {{ font-family: 'Microsoft YaHei', sans-serif; line-height: 1.6; color: #333; max-width: 800px; margin: 0 auto; }}
        h1 {{ font-size: 24pt; color: #2c3e50; text-align: center; border-bottom: 2px solid #3498db; margin: 30px 0; padding-bottom: 20px; }}
        h2 {{ font-size: 16pt; color: #2980b9; background-color: #ecf0f1; padding: 10px; border-left: 5px solid #2980b9; margin-top: 25px; }}
        table {{ width: 100%; border-collapse: collapse; margin-bottom: 20px; font-size: 10.5pt; }}
        th {{ background-color: #3498db; color: white; padding: 10px; text-align: left; border: 1px solid #2980b9; }}
        td {{ padding: 8px 10px; border: 1px solid #ddd; }}
        tr:nth-child(even) {{ background-color: #f9f9f9; }}
        .improvement-item {{ margin-bottom: 15px; padding-bottom: 15px; border-bottom: 1px dashed #eee; }}
        .improvement-title {{ font-weight: bold; color: #2c3e50; display: block; }}
        .summary-box {{ background-color: #fcf8e3; border: 1px solid #faebcc; padding: 15px; margin-bottom: 20px; }}
        .risk-box {{ border: 1px solid #ebccd1; background-color: #f2dede; color: #a94442; padding: 15px; }}
        .status-unclosed {{ color: #e74c3c; font-weight: bold; }}
        .status-closed {{ color: #27ae60; }}
    </style>
</head>
<body>
    <h1>{html.escape(project_name)} - 质量验收报告</h1>

    <h2>一、 验收概览</h2>
    <div class="summary-box">{summary_box}</div>
    <div class="risk-box">{top_risk}</div>

    <h2>二、 核心深度分析</h2>
    <table>
        <tr>
            <th>序号</th>
            <th>模块</th>
            <th>Bug数量</th>
            <th>权重</th>
            <th>模块密度(数量/权重)</th>
        </tr>
        {module_rows_html}
    </table>
    {module_interpretation_html}

    <h2>三、 改进措施</h2>
    <div>
        {improvements_html}
    </div>

    <h2>四、 附录：详细缺陷列表</h2>
    <table>
        <tr>
            <th>标题</th>
            <th>任务ID</th>
            <th>执行者</th>
            <th>父任务</th>
            <th>任务状态</th>
            <th>优先级</th>
        </tr>
        {appendix_rows_html}
    </table>
    {ai_section_html}
</body>
</html>
"""
    return html_doc


def normalize_priority_label(value):
    text = normalize_text(value)
    lowered = text.lower()
    if not text:
        return "未标注"
    if any(token in lowered for token in ["p0", "p1", "blocker", "critical"]) or any(
        token in text for token in ["高", "严重", "紧急"]
    ):
        return "高"
    if any(token in lowered for token in ["p2", "major", "medium"]) or "中" in text:
        return "中"
    if any(token in lowered for token in ["p3", "p4", "minor", "low"]) or "低" in text:
        return "低"
    return text


def _priority_sort_key(label):
    return {"高": 0, "中": 1, "低": 2, "未标注": 3}.get(label, 4), label


def collect_report_stats(rows):
    total = len(rows)
    resolved_count = sum(
        1 for row in rows if normalize_text(row.get("任务状态")) in CLOSED_STATUS_ORDER
    )
    high_count = sum(1 for row in rows if is_high_priority(row.get("优先级")))

    participants = set()
    module_counter = Counter()
    module_titles = defaultdict(list)
    priority_counter = Counter()
    executor_counter = Counter()
    status_counter = Counter()

    for row in rows:
        executor = normalize_text(row.get("执行者")) or "未分配"
        executor_counter[executor] += 1
        status_counter[normalize_text(row.get("任务状态")) or "未标注"] += 1
        priority_counter[normalize_priority_label(row.get("优先级"))] += 1

        for person in split_people(row.get("执行者")):
            participants.add(person)
        for person in split_people(row.get("参与者")):
            participants.add(person)

        title = normalize_text(row.get("标题"))
        module = classify_module(title)
        module_counter[module] += 1
        module_titles[module].append(title)

    module_analysis = []
    for module_name, count in module_counter.items():
        weight = density_weight(count)
        module_analysis.append(
            {
                "module": module_name,
                "count": count,
                "weight": weight,
                "density": count / weight,
                "percentage": format_percent(count, total),
            }
        )
    module_analysis.sort(key=lambda item: (-item["count"], item["module"]))

    top1_module = module_analysis[0]["module"] if module_analysis else "无"
    top1_nature = issue_nature_by_titles(module_titles.get(top1_module, []))
    unresolved_count = total - resolved_count

    return {
        "total": total,
        "resolved_count": resolved_count,
        "unresolved_count": unresolved_count,
        "high_count": high_count,
        "participants": participants,
        "module_analysis": module_analysis,
        "priority_counter": priority_counter,
        "executor_counter": executor_counter,
        "status_counter": status_counter,
        "top1_module": top1_module,
        "top1_nature": top1_nature,
    }


def _set_run_font(run, name=DOCX_BODY_FONT, size=10.5, bold=None, italic=None, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attribute), name)


def _configure_docx_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = DOCX_BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:ascii"), DOCX_BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), DOCX_BODY_FONT)

    heading_settings = {
        "Heading 1": (14, DOCX_ACCENT, False),
        "Heading 2": (13, DOCX_ACCENT, False),
        "Heading 3": (11.5, DOCX_ACCENT, False),
        "Heading 4": (10.5, DOCX_ACCENT, True),
    }
    for style_name, (size, color, italic) in heading_settings.items():
        style = document.styles[style_name]
        style.font.name = DOCX_HEADING_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = italic
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_HEADING_FONT)
        style._element.rPr.rFonts.set(qn("w:ascii"), DOCX_HEADING_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), DOCX_HEADING_FONT)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.space_before = Pt(10 if style_name != "Heading 1" else 24)
        style.paragraph_format.space_after = Pt(0)


def _set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def _set_cell_margins(cell, top=90, start=108, bottom=90, end=108):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _set_table_borders(table, color=DOCX_BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def _set_cell_text(cell, value, bold=False, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(normalize_text(value))
    _set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell)


def _add_docx_table(document, headers, rows, widths, font_size=10.5):
    table = document.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Light Grid Accent 1"
    except KeyError:
        table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    _mark_header_row(table.rows[0])

    for index, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Inches(widths[index])
        _set_cell_text(
            cell,
            header,
            bold=True,
            size=font_size,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        _set_cell_shading(cell, DOCX_HEADER_FILL)

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[index])
            align = (
                WD_ALIGN_PARAGRAPH.CENTER
                if index == 0 or (isinstance(value, (int, float)) and index < len(values) - 1)
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            _set_cell_text(cell, value, size=font_size, align=align)
            if row_index % 2 == 1:
                _set_cell_shading(cell, DOCX_ALT_ROW_FILL)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
            _set_cell_margins(cell)
    return table


def _add_docx_body_paragraph(document, text, bold_prefix=""):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        lead = paragraph.add_run(bold_prefix)
        _set_run_font(lead, bold=True)
        body = paragraph.add_run(text[len(bold_prefix):])
        _set_run_font(body)
    else:
        run = paragraph.add_run(text)
        _set_run_font(run)
    return paragraph


def _find_chart_font():
    candidates = [
        os.environ.get("REPORT_CJK_FONT", ""),
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf",
        "C:/Windows/Fonts/simsun.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate
    return ""


def _chart_font(size, bold=False):
    from PIL import ImageFont  # pylint: disable=import-outside-toplevel

    path = _find_chart_font()
    if path:
        try:
            return ImageFont.truetype(path, size=size, index=0)
        except (OSError, ValueError):
            pass
    return ImageFont.load_default()


def _chart_text_supports_cjk():
    path = _find_chart_font().lower()
    return any(token in path for token in ["msyh", "simhei", "simsun", "noto", "wqy"])


def _text_bbox(draw, text, font):
    left, top, right, bottom = draw.textbbox((0, 0), text, font=font)
    return right - left, bottom - top


def _make_priority_chart(priority_counter):
    try:
        from PIL import Image, ImageDraw  # pylint: disable=import-outside-toplevel
    except ImportError:
        return None

    width, height = 860, 470
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    total = sum(priority_counter.values())
    if total <= 0:
        return None

    colors = {"高": "#8BC66E", "中": "#F7C857", "低": "#5874C7"}
    ordered = sorted(priority_counter.items(), key=lambda item: _priority_sort_key(item[0]))
    box = (130, 45, 570, 405)
    start = -90
    for label, count in ordered:
        extent = count / total * 360
        draw.pieslice(box, start=start, end=start + extent, fill=colors.get(label, "#A9B4C6"), outline="white")
        start += extent
    inner = (225, 140, 475, 390)
    draw.ellipse(inner, fill="white")

    big_font = _chart_font(52)
    small_font = _chart_font(21)
    total_text = str(total)
    tw, th = _text_bbox(draw, total_text, big_font)
    draw.text(((inner[0] + inner[2] - tw) / 2, 205), total_text, fill="#333333", font=big_font)
    if _chart_text_supports_cjk():
        label = "Bug总数"
        tw, th = _text_bbox(draw, label, small_font)
        draw.text(((inner[0] + inner[2] - tw) / 2, 278), label, fill="#666666", font=small_font)

    legend_font = _chart_font(20)
    legend_x = 625
    legend_y = 105
    for label, count in ordered:
        color = colors.get(label, "#A9B4C6")
        draw.rounded_rectangle((legend_x, legend_y + 3, legend_x + 22, legend_y + 25), radius=3, fill=color)
        legend = f"{label}  {count}（{format_percent(count, total)}）" if _chart_text_supports_cjk() else f"{count} ({format_percent(count, total)})"
        draw.text((legend_x + 34, legend_y), legend, fill="#333333", font=legend_font)
        legend_y += 48

    stream = io.BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    return stream


def _make_module_chart(module_analysis):
    try:
        from PIL import Image, ImageDraw  # pylint: disable=import-outside-toplevel
    except ImportError:
        return None

    items = module_analysis[:8]
    if not items:
        return None
    width, height = 960, 520
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    left, top, right, bottom = 85, 45, 920, 425
    chart_height = bottom - top
    max_count = max(item["count"] for item in items) or 1
    font = _chart_font(17)
    value_font = _chart_font(18)
    label_font = _chart_font(15)
    bar_width = max(32, int((right - left) / len(items) * 0.62))
    step = (right - left) / len(items)
    grid_font = _chart_font(14)

    for tick in range(0, max_count + 1, max(1, max_count // 4 or 1)):
        y = bottom - int(tick / max_count * chart_height)
        draw.line((left, y, right, y), fill="#E1E5EB", width=1)
        draw.text((left - 35, y - 9), str(tick), fill="#555555", font=grid_font)
    draw.line((left, bottom, right, bottom), fill="#555555", width=2)
    draw.line((left, top, left, bottom), fill="#555555", width=2)

    for index, item in enumerate(items):
        x_center = int(left + step * (index + 0.5))
        bar_left = x_center - bar_width // 2
        bar_top = bottom - int(item["count"] / max_count * chart_height)
        draw.rectangle((bar_left, bar_top, x_center + bar_width // 2, bottom), fill="#7890D0")
        value = str(item["count"])
        vw, vh = _text_bbox(draw, value, value_font)
        draw.text((x_center - vw / 2, max(5, bar_top - vh - 4)), value, fill="#333333", font=value_font)
        if _chart_text_supports_cjk():
            label = item["module"]
            lw, lh = _text_bbox(draw, label, label_font)
            draw.text((x_center - lw / 2, bottom + 14), label, fill="#333333", font=label_font)
        else:
            draw.text((x_center - 5, bottom + 14), str(index + 1), fill="#333333", font=font)

    stream = io.BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    return stream


def _add_chart(document, stream, width_inches=5.75):
    if stream is None:
        return
    document.add_picture(stream, width=Inches(width_inches))
    paragraph = document.paragraphs[-1]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)


def _module_analysis_text(module_name, count, top_module):
    if module_name == top_module:
        return (
            f"问题主要集中在{module_name}，共 {count} 项，"
            "建议围绕核心流程、异常处理、结果交付和边界场景开展专项回归。"
        )
    return (
        f"{module_name} 共记录 {count} 项，"
        "建议结合缺陷标题逐项确认影响范围，避免同类问题在后续版本重复出现。"
    )


def _build_suggestion_groups(stats, ai_analysis):
    risks = ai_analysis.get("risks", []) if ai_analysis else []
    actions = ai_analysis.get("actions", []) if ai_analysis else []
    unresolved_text = f"当前未关闭缺陷 {stats['unresolved_count']} 项。"
    top_module = stats["top1_module"]
    top_count = stats["module_analysis"][0]["count"] if stats["module_analysis"] else 0

    defaults = [
        (
            "业务逻辑与结果链路",
            unresolved_text + f"问题较集中于{top_module}（{top_count} 项），需重点关注主流程闭环。",
            "围绕任务创建、处理、完成、失败和结果落地建立端到端回归路径，确保状态可追踪、失败可定位。",
        ),
        (
            "核心功能与技术稳定性",
            f"高优先级缺陷 {stats['high_count']} 项，部分问题可能影响核心功能稳定性。",
            "优先清理高优先级未闭环问题，补充异常重试、超时、并发和数据边界场景验证。",
        ),
        (
            "交互操作与兼容性",
            "缺陷标题中包含页面、显示、交互或提示相关问题时，容易在不同设备和状态下重复暴露。",
            "统一提示语、按钮状态和关键页面跳转规则，并补充多设备、多分辨率及异常状态回归。",
        ),
    ]
    if not ai_analysis:
        return defaults

    focus = [
        "发布风险与优先级",
        "核心模块与结果链路",
        "回归验证与体验",
    ]
    groups = []
    for index, (title, issue, suggestion) in enumerate(defaults):
        risk = risks[index] if index < len(risks) else ""
        action = actions[index] if index < len(actions) else ""
        if risk:
            issue = risk
        if action:
            suggestion = action
        groups.append((focus[index], issue, suggestion))
    return groups


def build_docx_bytes(project_name, rows, llm_config=None):
    stats = collect_report_stats(rows)
    ai_analysis = None
    if llm_config and llm_config.get("enabled"):
        try:
            ai_analysis = build_ai_insight(
                project_name=project_name,
                total=stats["total"],
                resolved_count=stats["resolved_count"],
                high_count=stats["high_count"],
                module_analysis=stats["module_analysis"],
                top1_module=stats["top1_module"],
                top1_nature=stats["top1_nature"],
                rows=rows,
                llm_config=llm_config,
            )
        except Exception as exc:  # pylint: disable=broad-except
            raise RuntimeError(f"AI 增强调用失败，报告未生成：{exc}") from exc

    document = Document()
    _configure_docx_styles(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    title = document.add_paragraph(style="Heading 1")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(16)
    title_run = title.add_run(f"{normalize_text(project_name)} Bug缺陷分享报告")
    _set_run_font(title_run, name=DOCX_HEADING_FONT, size=18, bold=True, color="000000")

    document.add_heading("一、版本信息", level=2)
    main_executors = "、".join(
        name
        for name, _ in stats["executor_counter"].most_common(3)
        if name != "未分配"
    ) or "未分配"
    version_rows = [
        ("版本号", normalize_text(project_name)),
        ("测试日期", dt.datetime.now().strftime("%Y年%m月%d日")),
        ("Bug总数", f"{stats['total']} 个"),
        (
            "已修复缺陷",
            f"{stats['resolved_count']} 个（修复率 {format_percent(stats['resolved_count'], stats['total'])}）",
        ),
        ("涉及功能模块", f"{len(stats['module_analysis'])} 个"),
        ("主要执行人", main_executors),
    ]
    _add_docx_table(document, ["项目指标", "统计结果"], version_rows, [1.45, 4.55])

    document.add_heading("二、Bug优先级分布", level=2)
    _add_chart(document, _make_priority_chart(stats["priority_counter"]))
    priority_rows = [
        (label, count, format_percent(count, stats["total"]))
        for label, count in sorted(stats["priority_counter"].items(), key=lambda item: _priority_sort_key(item[0]))
    ]
    _add_docx_table(document, ["优先级", "数量", "占比"], priority_rows, [2.0, 2.0, 2.0])

    document.add_heading("三、Bug功能模块分布", level=2)
    _add_chart(document, _make_module_chart(stats["module_analysis"]))
    module_rows = [
        (item["module"], item["count"], item["percentage"])
        for item in stats["module_analysis"]
    ]
    _add_docx_table(document, ["功能模块", "数量", "占比"], module_rows, [2.6, 1.6, 1.8])

    document.add_heading("四、Bug执行人分布", level=2)
    executor_rows = [
        (name, count, format_percent(count, stats["total"]))
        for name, count in stats["executor_counter"].most_common()
    ]
    _add_docx_table(document, ["执行人", "处理数量", "占比"], executor_rows, [2.6, 1.6, 1.8])

    document.add_heading("五、测试结论与建议", level=2)
    document.add_heading("5.1 测试结论", level=3)
    conclusion_table = [
        ("缺陷总数", f"{stats['total']} 个"),
        (
            "已修复缺陷",
            f"{stats['resolved_count']} 个（修复率 {format_percent(stats['resolved_count'], stats['total'])}）",
        ),
        ("未关闭缺陷", f"{stats['unresolved_count']} 个"),
        ("高优先级缺陷", f"{stats['high_count']} 个"),
        (
            "问题集中模块",
            (
                f"{stats['top1_module']}（{stats['module_analysis'][0]['count']} 个）"
                if stats["module_analysis"]
                else "无"
            ),
        ),
    ]
    _add_docx_table(document, ["评估项", "评估结果"], conclusion_table, [1.45, 4.55])

    if ai_analysis:
        decision_text = {
            "GO": "具备发布条件",
            "NO-GO": "暂不建议发布",
            "CONDITIONAL-GO": "满足风险收敛条件后发布",
        }.get(ai_analysis["decision"], ai_analysis["decision"])
        conclusion_text = (
            f"{ai_analysis['summary']} 当前模型发布建议：{decision_text}。"
        )
    elif stats["unresolved_count"] == 0:
        conclusion_text = (
            "本轮缺陷已全部关闭，版本整体质量趋于稳定，具备发布条件。"
            "建议上线后继续关注高频使用路径和真实用户场景。"
        )
    else:
        conclusion_text = (
            f"本轮共记录缺陷 {stats['total']} 项，已关闭 {stats['resolved_count']} 项，"
            f"仍有 {stats['unresolved_count']} 项未关闭。问题主要集中在"
            f"{stats['top1_module']}，建议在发布前完成高优先级缺陷闭环并补充专项回归。"
        )
    _add_docx_body_paragraph(document, "综合评估：" + conclusion_text, bold_prefix="综合评估：")

    document.add_heading("5.2 缺陷问题分析", level=3)
    top_modules = "、".join(
        f"{item['module']}（{item['count']}项）"
        for item in stats["module_analysis"][:3]
    ) or "暂无模块数据"
    analysis_text = (
        f"本轮共记录缺陷 {stats['total']} 项，问题主要集中在 {top_modules}。"
        f"其中高优先级缺陷 {stats['high_count']} 项，未关闭缺陷 {stats['unresolved_count']} 项，"
        "说明当前版本仍需关注核心流程稳定性、异常兜底和交互一致性。"
    )
    if ai_analysis and ai_analysis["risks"]:
        analysis_text += " 模型识别的主要风险包括：" + "；".join(ai_analysis["risks"]) + "。"
    _add_docx_body_paragraph(document, analysis_text)

    problem_rows = [
        (
            item["module"],
            f"{item['count']} 个",
            _module_analysis_text(item["module"], item["count"], stats["top1_module"]),
        )
        for item in stats["module_analysis"][:6]
    ]
    _add_docx_table(document, ["问题类型", "Bug数量", "问题分析"], problem_rows, [1.5, 0.9, 3.6], font_size=9.5)

    document.add_heading("5.3 优化建议与风险提示", level=3)
    _add_docx_body_paragraph(
        document,
        "结合本轮缺陷分布与问题表现，建议从业务逻辑、功能稳定性及交互体验三个层面同步推进优化，以降低重复性问题并提升版本交付质量。",
    )
    for heading, issue, suggestion in _build_suggestion_groups(stats, ai_analysis):
        document.add_heading(heading, level=4)
        _add_docx_table(
            document,
            ["关注点", "问题表现", "优化建议"],
            [(heading, issue, suggestion)],
            [1.45, 2.15, 2.4],
            font_size=9.5,
        )

    document.add_page_break()
    document.add_heading("六、Bug详细列表", level=2)
    detail_rows = []
    for index, row in enumerate(sort_rows_for_appendix(rows), start=1):
        task_id = normalize_text(row.get("任务ID"))
        title_text = normalize_text(row.get("标题")) or "未填写标题"
        if task_id:
            title_text = f"[{task_id}] {title_text}"
        detail_rows.append(
            (
                index,
                title_text,
                normalize_priority_label(row.get("优先级")),
                normalize_text(row.get("执行者")) or "未分配",
                normalize_text(row.get("任务状态")) or "未标注",
            )
        )
    _add_docx_table(
        document,
        ["序号", "Bug标题", "优先级", "执行人", "状态"],
        detail_rows,
        [0.45, 3.45, 0.7, 0.75, 0.65],
        font_size=9,
    )

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def save_report(input_file, project_name):
    headers, rows = load_rows(input_file)
    validate_columns(headers)
    rendered = build_docx_bytes(project_name, rows)

    output_dir = ensure_f_drive_or_desktop()
    timestamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    output_name = f"Bug缺陷分享报告_{timestamp}.docx"
    output_path = output_dir / output_name

    output_path.write_bytes(rendered)
    return output_path


def main():
    from tkinter import Tk, filedialog, messagebox, simpledialog  # pylint: disable=import-outside-toplevel

    root = Tk()
    root.withdraw()
    root.update()

    file_path = filedialog.askopenfilename(
        title="请选择 Bug 数据文件 (CSV/Excel)",
        filetypes=[("Data Files", "*.csv *.xlsx *.xls"), ("All Files", "*.*")],
    )
    if not file_path:
        messagebox.showinfo("提示", "未选择文件，已取消。")
        return

    default_project = Path(file_path).stem
    project_name = simpledialog.askstring("项目名称", "请输入项目名称：", initialvalue=default_project)
    if not project_name:
        project_name = default_project

    try:
        output_path = save_report(file_path, project_name)
        target_tip = "F盘" if str(output_path).lower().startswith("f:") else "桌面(未检测到F盘)"
        messagebox.showinfo("生成成功", f"报告已生成到 {target_tip}\n{output_path}")
        print(f"生成成功: {output_path}")
    except Exception as exc:  # pylint: disable=broad-except
        messagebox.showerror("生成失败", str(exc))
        print(f"生成失败: {exc}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
