# -*- coding: utf-8 -*-

import io
import unittest
import zipfile

from docx import Document

from generate_bug_report import build_docx_bytes


class DocxReportTests(unittest.TestCase):
    def test_builds_reference_style_report(self):
        rows = [
            {
                "标题": "视频翻译导出失败",
                "任务ID": "BUG-1",
                "执行者": "测试一",
                "父任务": "",
                "任务状态": "已解决",
                "优先级": "高",
                "参与者": "测试一",
            },
            {
                "标题": "AI字幕页面按钮显示异常",
                "任务ID": "BUG-2",
                "执行者": "测试二",
                "父任务": "",
                "任务状态": "待处理",
                "优先级": "中",
                "参与者": "测试二",
            },
            {
                "标题": "多语言提示语问题",
                "任务ID": "BUG-3",
                "执行者": "测试一",
                "父任务": "",
                "任务状态": "已完成",
                "优先级": "低",
                "参与者": "测试一",
            },
        ]

        content = build_docx_bytes("示例版本", rows)
        document = Document(io.BytesIO(content))

        headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
        self.assertIn("一、版本信息", headings)
        self.assertIn("五、测试结论与建议", headings)
        self.assertIn("六、Bug详细列表", headings)
        self.assertEqual(len(document.tables), 10)
        self.assertEqual(len(document.inline_shapes), 2)
        self.assertEqual(len(document.tables[-1].rows), 4)

    def test_output_is_a_valid_docx_package(self):
        content = build_docx_bytes("空数据版本", [])
        with zipfile.ZipFile(io.BytesIO(content)) as package:
            self.assertIn("word/document.xml", package.namelist())
            self.assertEqual(
                len([name for name in package.namelist() if name.startswith("word/media/")]),
                0,
            )


if __name__ == "__main__":
    unittest.main()
